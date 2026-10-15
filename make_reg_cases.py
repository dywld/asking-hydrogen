#!/usr/bin/env python3
"""규제 애로 사례 — 수소경제 포털(h2hub)의 '규제애로 개선현황'과 '규제개선 건의'.

사업자 질문이 이미 알려진 규제 애로와 겹치면 답 아래에 사례를 붙인다
("이건 지금 규제로 막혀 있고, 이렇게 개선됐다/건의 중이다").

    python make_reg_cases.py       # → data/reg_cases.json
"""
import json
import re
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
SRC = HERE / "문서" / "수소사업자_질문애로_수집.xlsx"
OUT = HERE / "data" / "reg_cases.json"


def clean(s, n=220):
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    s = re.sub(r"^(개선건의\s*-\s*개선할 내용 명확히 작성\s*-|-)\s*", "", s)   # 건의 양식의 머리말
    return s[:n]


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True)
    out = []
    for r in list(wb["규제애로_개선현황"].iter_rows(min_row=2, values_only=True)):
        field, title, dept, pain, state, fix = (r + (None,) * 6)[:6]
        if title:
            out.append({"kind": "개선현황", "field": clean(field, 20), "title": clean(title, 80),
                        "dept": clean(dept, 30), "pain": clean(pain), "state": clean(state, 60),
                        "fix": clean(fix), "law": "", "link": "https://h2hub.or.kr"})
    for r in list(wb["규제개선_건의"].iter_rows(min_row=2, values_only=True)):
        field, title, dept, law, pain, fix, link = (r + (None,) * 7)[:7]
        if title:
            out.append({"kind": "건의", "field": clean(field, 20), "title": clean(title, 80),
                        "dept": clean(dept, 30), "pain": clean(pain), "state": "건의(처리 결과 미확인)",
                        "fix": clean(fix), "law": clean(law, 80), "link": link or "https://h2hub.or.kr"})
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(out)}건 → {OUT.name}")


if __name__ == "__main__":
    main()

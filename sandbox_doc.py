#!/usr/bin/env python3
"""규제샌드박스 신청 서류 만들기 — 공식 신청서 PDF 채우기 + 사업실시계획서 워드.

신청서는 법령정보센터의 원본 서식(산업융합촉진법 시행령 별지 제3호·제7호)을 받아
칸 이름 옆에 글을 얹는다. 칸 위치는 좌표를 박지 않고 칸 이름을 찾아서 정한다
(서식이 개정돼 칸이 조금 움직여도 따라가게).
사업자만 아는 칸은 주황색 '[직접 입력]'으로 남긴다. 지어낸 값은 넣지 않는다.
"""
import io
import re
import time

import forms

FORM_NO = {"실증특례": 3, "임시허가": 7}
LAW = "산업융합촉진법 시행령"
FONT = "C:/Windows/Fonts/malgun.ttf"
BLUE = (0.12, 0.3, 0.75)
ORANGE = (0.75, 0.35, 0.0)
NEED = "[직접 입력]"


def _form_pdf(kind):
    r = forms.resolve(LAW, "서식", FORM_NO[kind])
    if not r or not r.get("_pdf"):
        raise RuntimeError("공식 서식 파일을 찾지 못했습니다")
    data, _ = forms.fetch_file(r["_pdf"], "pdf")
    return data


# 칸: (key, 칸 이름, 질문, 왼쪽 칸 이름만 볼지, 여러 줄 높이). 칸 이름은 PDF 에서 찾을 글자 그대로.
SLOTS = [
    ("company", "이름(회사명)", "회사 이름 ('이름(회사명)' 칸)", True, 0),
    ("bizno", "사업자(법인)등록번호", "사업자(법인)등록번호", False, 0),
    ("address", "주소", "회사 주소 ('주소' 칸)", False, 0),
    ("ceo", "대표자", "대표자 이름 ('대표자' 칸)", False, 0),
    ("tel", "전화번호", "전화번호", False, 0),
    ("email", "전자우편", "전자우편", False, 0),
    ("name", "명칭", "신기술·서비스 이름 ('명칭' 칸)", True, 0),
    ("what", "주요내용", "주요 내용 ('주요내용' 칸)", True, 26),
    ("laws", "실증을 위한 규제특례 법령", "특례가 필요한 법령", False, 20),
    ("area", "구역", "실증 구역", True, 0),
    ("period", "기간", "실증 기간", True, 0),
    ("scale", "규모", "실증 규모", True, 0),
    ("applicant", "신청인 성명", "신청인 이름 ('신청인 성명' 칸)", False, 0),
]


def _hit(p, label, first_left):
    hits = p.search_for(label)
    if first_left:
        hits = [h for h in hits if h.x0 < 200]
    return hits[0] if hits else None


# 유형·신청사유의 '[ ]' 칸. 유형은 하나만, 신청사유는 여러 개 고를 수 있다. 값은 고른 순번 "0,2".
CHK_Q = {"type": ["신제품인 경우", "서비스인 경우", "신제품과 서비스가 융합된 경우"],
         "why": ["맞는 기준·규격·요건이 없는 경우", "기준·규격·요건을 적용하는 것이 맞지 않는 경우",
                 "제한된 구역·기간·규모 안에서 실증이 필요한 경우"]}


def _checks(p):
    """→ {"type": [Rect…], "why": [Rect…]}  각 '[' 글자 자리(위에서 아래 순서)."""
    hs = sorted(p.search_for("["), key=lambda h: h.y0)
    ty, sa = p.search_for("유형"), p.search_for("신청사유")
    out = {"type": [], "why": []}
    if ty:
        out["type"] = [h for h in hs if 170 < h.x0 < 230 and abs(h.y0 - ty[0].y0) < 30]
    if sa:
        out["why"] = [h for h in hs if 120 < h.x0 < 160 and abs(h.y0 - sa[0].y0) < 70]
    return out


def picked(f, key):
    return {int(x) for x in re.findall(r"\d", str(f.get(key) or ""))}


def _dyn(p):
    """날짜(년·월·일 자리)와 서명 자리('(서명 또는 인)' 위)."""
    out = {"date": [], "sign": None}
    d = p.search_for("년     월     일")
    if d:
        r = d[0]
        for key, x in (("date_y", r.x0 - 24), ("date_m", r.x0 + 15), ("date_d", r.x0 + 43)):
            out["date"].append((key, x, r.y1 - 1))
    sg = p.search_for("(서명 또는 인)")
    if sg:
        g = sg[-1]
        top = max(g.y0 - 12, d[0].y1 + 3) if d else g.y0 - 12
        out["sign"] = (g.x0 - 6, top, g.x1 + 6, g.y1 + 8)
    return out


def fill_pdf(f, ink="black", dynamic=True):
    """f: kind·company·name·what·area·period·scale·laws (+ bizno·address·ceo·tel·email·applicant·
    date_y/m/d·sign 은 화면에서 바로 고친 값).
    ink="black"(제출용): 검정 글씨, 빈칸은 비운다. "blue"(검토용): 파랑, 빈칸은 '[직접 입력]'.
    dynamic=False 면 날짜·서명을 비운다(미리보기에서 바로 고칠 때 바탕용)."""
    import pymupdf
    kind = f.get("kind") if f.get("kind") in FORM_NO else "실증특례"
    doc = pymupdf.open(stream=_form_pdf(kind), filetype="pdf")
    p = doc[0]
    p.insert_font(fontname="mg", fontfile=FONT)
    right = 532
    blue = ink == "blue"
    INK = BLUE if blue else (0, 0, 0)

    for key, label, _, first_left, box_h in SLOTS:
        h = _hit(p, label, first_left)
        if not h:
            continue
        value = (f.get(key) or "").strip()
        if not value and not blue:
            continue
        text, color = (value, INK) if value else (NEED, ORANGE)
        if box_h:
            p.insert_textbox(pymupdf.Rect(h.x1 + 6, h.y0 - 1, right, h.y0 + box_h), text,
                             fontname="mg", fontsize=7.5, color=color, lineheight=1.25)
        else:
            p.insert_text((h.x1 + 6, h.y1 - 1), text[:60], fontname="mg", fontsize=8, color=color)
    for key, rects in _checks(p).items():                # 화면에서 고른 유형·신청사유에 √
        for i, h in enumerate(rects):
            if i in picked(f, key):
                p.insert_text((h.x0 + 4.5, h.y1 - 1.5), "√", fontname="mg", fontsize=9, color=INK)
    if dynamic:
        lay = _dyn(p)
        today = {"date_y": time.strftime("%Y"), "date_m": time.strftime("%m").lstrip("0"),
                 "date_d": time.strftime("%d").lstrip("0")}
        for key, x, y in lay["date"]:
            p.insert_text((x, y), (f.get(key) or today[key]).strip()[:4], fontname="mg", fontsize=8, color=INK)
        sig = f.get("sign") or ""
        if lay["sign"] and sig.startswith("data:image/png;base64,"):
            import base64
            # 잉크 부분만 잘라 온 그림을 원래 비율 그대로 '(서명 또는 인)' 위에 겹친다(한국 서식 관례).
            p.insert_image(pymupdf.Rect(*lay["sign"]), stream=base64.b64decode(sig.split(",", 1)[1]),
                           keep_proportion=True, overlay=True)
    # 유형·신청사유는 사업자가 판단할 일이라 AI 가 고르지 않는다(미리보기에서 눌러 고른다).
    note = "※ 유형·신청사유는 미리보기에서 눌러 고르세요. 파란 글씨는 AI 초안이니 원문과 대조하세요."
    if blue:
        p.insert_text((40, 830), note, fontname="mg", fontsize=7, color=ORANGE)
    doc.subset_fonts()          # 한글 글꼴을 통째로 넣으면 13MB가 된다(실측). 쓴 글자만 남긴다.
    return doc.tobytes(garbage=3, deflate=True)


def edit_layout(f):
    """미리보기에서 바로 고치기용(허가 신청서 permit_forms.edit_layout 과 같은 방식).
    칸 글자·날짜·서명을 비운 바탕 PDF 와, 칸마다 입력 상자를 얹을 자리(pt)를 돌려준다."""
    import pymupdf
    kind = f.get("kind") if f.get("kind") in FORM_NO else "실증특례"
    data = fill_pdf({"kind": kind}, "black", dynamic=False)
    doc = pymupdf.open(stream=data, filetype="pdf")
    p = doc[0]
    pw, ph = p.rect.width, p.rect.height
    words = p.get_text("words")
    vlines = []
    for d in p.get_drawings():
        for it in d["items"]:
            if it[0] == "l" and abs(it[1].x - it[2].x) < 1 and abs(it[1].y - it[2].y) > 5:
                vlines.append((it[1].x, min(it[1].y, it[2].y), max(it[1].y, it[2].y)))
            elif it[0] == "re" and it[1].height > 5:
                vlines += [(it[1].x0, it[1].y0, it[1].y1), (it[1].x1, it[1].y0, it[1].y1)]
    boxes = []
    for key, label, q, first_left, box_h in SLOTS:
        h = _hit(p, label, first_left)
        if not h:
            continue
        x = h.x1 + 6
        top, bottom = (h.y0 - 1, h.y0 + box_h) if box_h else (h.y1 - 10, h.y1 + 2)
        right = 532
        for w in words:                          # 같은 줄 다음 글자 직전까지
            if w[0] > x + 6 and w[1] < bottom and w[3] > top:
                right = min(right, w[0] - 4)
        for vx, y0, y1 in vlines:                # 표 세로줄을 넘지 않게
            if vx > x + 6 and y0 < bottom and y1 > top:
                right = min(right, vx - 4)
        boxes.append({"key": key, "q": q, "x": round(x, 1), "y": round(top, 1),
                      "w": round(max(30, right - x), 1), "h": round(bottom - top, 1),
                      "fs": 7.5 if box_h else 8, "value": (f.get(key) or "").strip()})
    for key, rects in _checks(p).items():                # 줄 전체를 누르는 자리(√ 는 '[' 안에 보인다)
        for i, h in enumerate(rects):
            nxt = rects[i + 1].y0 if i + 1 < len(rects) else h.y1 + (h.y0 - rects[i - 1].y0 if i else 12) - h.height
            boxes.append({"key": key, "kind": "check", "i": i, "one": key == "type", "q": CHK_Q[key][i] if i < len(CHK_Q[key]) else "",
                          "x": round(h.x0 - 2, 1), "y": round(h.y0 - 1, 1), "w": round(530 - h.x0, 1),
                          "h": round(max(10, nxt - h.y0 - 1), 1), "fs": 9, "value": str(f.get(key) or "")})
    lay = _dyn(p)
    today = {"date_y": time.strftime("%Y"), "date_m": time.strftime("%m").lstrip("0"),
             "date_d": time.strftime("%d").lstrip("0")}
    for key, x, y in lay["date"]:
        w = 24 if key == "date_y" else 13
        boxes.append({"key": key, "q": {"date_y": "신청 연도", "date_m": "월", "date_d": "일"}[key],
                      "x": round(x - 1, 1), "y": round(y - 9, 1), "w": w, "h": 12, "fs": 8,
                      "value": (f.get(key) or today[key]).strip()})
    if lay["sign"]:
        x0, y0, x1, y1 = lay["sign"]
        boxes.append({"key": "sign", "kind": "sign", "q": "서명", "x": round(x0, 1), "y": round(y0, 1),
                      "w": round(x1 - x0, 1), "h": round(y1 - y0, 1), "fs": 8, "value": ""})
    return data, boxes, pw, ph


_BOLD = re.compile(r"\*\*(.+?)\*\*")
_CITE = re.compile(r"\[(\d+(?:\s*,\s*\d+)*)\]")


def make_docx(f, text, refs):
    """사업실시계획서 초안. text: 답변(마크다운), refs: {'1': '법령 조문', ...}."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Malgun Gothic"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    st.font.size = Pt(10.5)
    kind = f.get("kind") if f.get("kind") in FORM_NO else "실증특례"
    doc.add_heading(f"사업실시계획서 ({kind})", level=0)
    p = doc.add_paragraph(f"산업융합촉진법 제10조의{'3' if kind == '실증특례' else '6'}제6항제1호 · "
                          f"신청서 별지 제{FORM_NO[kind]}호서식 붙임서류 1")
    p.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, label in (("company", "신청인"), ("name", "명칭"), ("what", "주요 내용"),
                     ("area", "구역"), ("period", "기간"), ("scale", "규모")):
        row = t.add_row().cells
        row[0].text = label
        row[1].text = f.get(k) or NEED

    def cite(m):
        names = [refs.get(n.strip()) for n in m.group(1).split(",")]
        names = [n for n in names if n]
        return f" (근거: {'; '.join(names)})" if names else ""

    for raw in (text or "").splitlines():
        line = raw.rstrip()
        if not line.strip() or line.strip() == "---":
            continue
        line = _CITE.sub(cite, line)
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            doc.add_heading(_BOLD.sub(r"\1", m.group(2)), level=min(len(m.group(1)), 3))
            continue
        bullet = re.match(r"^\s*[-*•]\s+(.*)", line)
        body = bullet.group(1) if bullet else line
        para = doc.add_paragraph(style="List Bullet" if bullet else None)
        # **굵게** 는 살리고, [직접 입력 …] 은 주황으로
        for part in re.split(r"(\*\*.+?\*\*|\[직접 입력[^\]]*\])", body):
            if not part:
                continue
            run = para.add_run(part.strip("*") if part.startswith("**") else part)
            if part.startswith("**"):
                run.bold = True
            if part.startswith("[직접 입력"):
                run.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)
    # 붙임서류 2번부터는 작성 틀만 준다(사업자만 아는 내용이라 AI가 지어내지 않는다). 시행령 제11조의3·제11조의6 순서.
    doc.add_page_break()
    doc.add_heading("붙임서류 작성 틀", level=1)
    name = f.get("name") or "[직접 입력: 제품·서비스 이름]"
    tpl = [
        ("산업융합 신제품·서비스 해당 설명", [f"{name}의 핵심 기술·구성", "기존 제품·서비스와 다른 점", "융합된 기술·분야"]),
        ("현행 법령으로 안 되는 사유", ["막히는 법령·조문: " + (f.get("laws") or "[직접 입력]"),
                                 "사유: 기준 없음 / 기준이 맞지 않음 / 불합리 중 해당", "필요한 특례 내용"]),
        ("혁신성 및 이용자 편익", ["혁신성(새로운 점)", "이용자·국민이 얻는 이득", "산업·지역 파급효과"]),
    ] + ([("향후 시장 성장 가능성", ["시장 규모·성장률(출처)", "수요처·고객", "사업화 계획"])] if kind == "실증특례" else []) + [
        ("손해 및 손해배상 방안", ["가입할 책임보험(보험사·보상한도)", "사고 시 배상 절차"]),
        ("생명·건강·안전·환경·개인정보 영향", ["위험요소와 안전대책(KGS Code 등 근거)", "시험·인증 결과",
                                       "환경·개인정보 영향과 대책"]),
    ]
    for i, (title, items) in enumerate(tpl, 2):
        doc.add_heading(f"붙임 {i}. {title}", level=2)
        for it in items:
            para = doc.add_paragraph(style="List Bullet")
            para.add_run(it + " : ")
            r = para.add_run("[직접 입력]")
            r.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)
    end = doc.add_paragraph("이 문서는 AI가 만든 초안입니다. 근거 조문은 원문과 대조하고, "
                            "[직접 입력] 칸을 채운 뒤 제출하세요. 계획서 세부 양식은 KIAT 안내를 확인하세요.")
    end.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_hwpx(f, text):
    """사업실시계획서 초안을 한글(HWPX)로. make_docx 와 같은 내용(표 + 본문), 인용 번호는 뺀다."""
    import hwpx_doc
    kind = f.get("kind") if f.get("kind") in FORM_NO else "실증특례"
    blocks = [
        ("title", f"사업실시계획서 ({kind})"),
        ("para", f"산업융합촉진법 제10조의{'3' if kind == '실증특례' else '6'}제6항제1호 · "
                 f"신청서 별지 제{FORM_NO[kind]}호서식 붙임서류 1"),
        ("table", [(label, f.get(k) or NEED) for k, label in
                   (("company", "신청인"), ("name", "명칭"), ("what", "주요 내용"),
                    ("area", "구역"), ("period", "기간"), ("scale", "규모"))]),
    ]
    for raw in (text or "").splitlines():
        line = _CITE.sub("", raw.rstrip())
        if not line.strip() or line.strip() == "---":
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            blocks.append(("heading", _BOLD.sub(r"\1", m.group(2))))
            continue
        bullet = re.match(r"^\s*[-*•]\s+(.*)", line)
        body = _BOLD.sub(r"\1", bullet.group(1) if bullet else line)
        blocks.append(("para", ("· " if bullet else "") + body))
    blocks.append(("para", "이 문서는 AI가 만든 초안입니다. 근거 조문은 원문과 대조하고, "
                           "[직접 입력] 칸을 채운 뒤 제출하세요. 계획서 세부 양식은 KIAT 안내를 확인하세요."))
    return hwpx_doc.make_hwpx(blocks, title=f"사업실시계획서 ({kind})")
